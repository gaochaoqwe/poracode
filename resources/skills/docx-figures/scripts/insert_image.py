#!/usr/bin/env python3
"""Insert image(s) (optionally with captions) into an existing .docx.

Single-image mode finds the first body paragraph containing `--anchor`, then
inserts a picture paragraph (and an optional caption paragraph) before or after
it, and saves.

Batch mode (`--manifest <json>`) opens the document once, applies every figure
entry in array order (captions auto-number continuously across entries), and
saves once -- so N images cost one parse and one save instead of N. It is
all-or-nothing: nothing is written unless every entry succeeds.

Built on python-docx. Kept narrow so the agent never hand-writes docx code.

Example (single):
  python insert_image.py --docx out.docx --image fig1.png \
      --anchor "3.2 软件能力需求" --caption "图3-1 系统功能组成" --width-cm 14

Example (batch) -- manifest.json:
  {
    "docx": "docs/out-filled.docx",
    "out": "docs/out-figures.docx",
    "figures": [
      {"image": "docs/figures/fig1.png", "anchor": "本软件是…的核心组成软件",
       "chapter": 3, "caption_title": "系统功能组成", "width_cm": 14}
    ]
  }
  python insert_image.py --manifest manifest.json
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.text.paragraph import Paragraph

FIGURE_CAPTION_RE = re.compile(r"^图\s*(\d+)\s*[-－]\s*(\d+)")


class InsertError(Exception):
    """A recoverable input problem (bad anchor, bad caption args), so a batch
    loop can report which entry failed instead of dumping a traceback."""


def next_figure_number(doc, chapter: int) -> int:
    """Highest existing `图<chapter>-N` index + 1 (1 when none exist)."""
    highest = 0
    for paragraph in doc.paragraphs:
        match = FIGURE_CAPTION_RE.match(paragraph.text.strip())
        if match and int(match.group(1)) == chapter:
            highest = max(highest, int(match.group(2)))
    return highest + 1


def find_paragraph(doc, anchor: str, occurrence: int):
    matches = [p for p in doc.paragraphs if anchor in p.text]
    if len(matches) < occurrence or occurrence < 1:
        raise InsertError(
            f"anchor not found (occurrence {occurrence}): {anchor!r} "
            f"(matched {len(matches)})"
        )
    return matches[occurrence - 1]


def make_paragraph_after(reference):
    element = OxmlElement("w:p")
    reference.addnext(element)
    return element


def insert_one(
    doc,
    *,
    image: str,
    anchor: str,
    caption: str = "",
    chapter: int = 0,
    caption_title: str = "",
    width_cm: float = 14.0,
    position: str = "after",
    occurrence: int = 1,
) -> None:
    """Insert one picture (+ optional caption) into an already-open `doc`.

    Mutates `doc` in memory only; the caller owns saving. Raises InsertError on
    bad input so a batch loop can report the failing entry by index. Caption
    auto-numbering re-scans `doc` each call, so entries inserted earlier in the
    same batch are counted and numbering stays continuous.
    """
    if position not in ("after", "before"):
        raise InsertError(f"position must be 'after' or 'before', got {position!r}")
    target = find_paragraph(doc, anchor, occurrence)

    if position == "after":
        image_element = make_paragraph_after(target._p)
    else:
        image_element = OxmlElement("w:p")
        target._p.addprevious(image_element)

    image_para = Paragraph(image_element, target._parent)
    image_para.alignment = 1  # center
    image_para.add_run().add_picture(image, width=Cm(width_cm))

    caption_text = caption
    if not caption_text and caption_title:
        if chapter <= 0:
            raise InsertError("caption_title requires a positive chapter")
        caption_text = f"图{chapter}-{next_figure_number(doc, chapter)} {caption_title}"

    if caption_text:
        caption_element = OxmlElement("w:p")
        image_element.addnext(caption_element)
        caption_para = Paragraph(caption_element, target._parent)
        caption_para.alignment = 1
        run = caption_para.add_run(caption_text)
        try:
            caption_para.style = doc.styles["Caption"]
        except KeyError:
            run.italic = True


def save_doc(doc, out: str, source: str) -> None:
    """Save `doc` to `out`.

    When overwriting the source in place, write to a temp sibling then
    atomically replace, so a crash mid-serialize cannot truncate the original.
    """
    parent = os.path.dirname(os.path.abspath(out))
    if parent:
        os.makedirs(parent, exist_ok=True)
    if os.path.abspath(out) == os.path.abspath(source):
        tmp = out + ".tmp"
        doc.save(tmp)
        os.replace(tmp, out)
    else:
        doc.save(out)


def run_single(args) -> int:
    for name in ("docx", "image", "anchor"):
        if not getattr(args, name):
            raise SystemExit(f"--{name} is required (or use --manifest)")
    doc = Document(args.docx)
    try:
        insert_one(
            doc,
            image=args.image,
            anchor=args.anchor,
            caption=args.caption,
            chapter=args.chapter,
            caption_title=args.caption_title,
            width_cm=args.width_cm,
            position=args.position,
            occurrence=args.occurrence,
        )
    except InsertError as exc:
        raise SystemExit(str(exc)) from exc
    out = args.out or args.docx
    save_doc(doc, out, args.docx)
    print(f"image inserted: {out}")
    return 0


def run_manifest(args) -> int:
    try:
        with open(args.manifest, "r", encoding="utf-8") as handle:
            manifest = json.load(handle)
    except OSError as exc:
        raise SystemExit(f"cannot read manifest: {exc}") from exc
    except json.JSONDecodeError as exc:
        raise SystemExit(f"manifest is not valid JSON: {exc}") from exc

    if not isinstance(manifest, dict):
        raise SystemExit("manifest must be a JSON object")
    docx = manifest.get("docx")
    figures = manifest.get("figures")
    if not isinstance(docx, str) or not docx:
        raise SystemExit("manifest requires a string 'docx'")
    if not isinstance(figures, list) or not figures:
        raise SystemExit("manifest requires a non-empty 'figures' list")

    out = manifest.get("out") or docx
    doc = Document(docx)

    for index, item in enumerate(figures):
        if not isinstance(item, dict):
            raise SystemExit(f"manifest figure[{index}] must be an object")
        image = item.get("image")
        anchor = item.get("anchor")
        if not isinstance(image, str) or not image or not isinstance(anchor, str) or not anchor:
            raise SystemExit(f"manifest figure[{index}] requires 'image' and 'anchor'")
        try:
            insert_one(
                doc,
                image=image,
                anchor=anchor,
                caption=item.get("caption", ""),
                chapter=item.get("chapter", 0),
                caption_title=item.get("caption_title", ""),
                width_cm=item.get("width_cm", 14.0),
                position=item.get("position", "after"),
                occurrence=item.get("occurrence", 1),
            )
        except Exception as exc:  # noqa: BLE001 - report the failing entry index
            raise SystemExit(f"manifest figure[{index}] failed: {exc}") from exc

    # Reached only when every entry succeeded -- all-or-nothing.
    save_doc(doc, out, docx)
    print(f"{len(figures)} images inserted: {out}")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Insert image(s) into a .docx")
    parser.add_argument("--manifest", default="", help="Batch: JSON manifest of figures")
    parser.add_argument("--docx", default="")
    parser.add_argument("--image", default="")
    parser.add_argument("--anchor", default="", help="Substring of the target paragraph")
    parser.add_argument("--caption", default="", help="Exact caption text")
    parser.add_argument(
        "--chapter",
        type=int,
        default=0,
        help="With --caption-title, auto-numbers as 图<chapter>-<next>",
    )
    parser.add_argument("--caption-title", default="")
    parser.add_argument("--width-cm", type=float, default=14.0)
    parser.add_argument("--position", choices=["after", "before"], default="after")
    parser.add_argument("--occurrence", type=int, default=1)
    parser.add_argument("--out", default="", help="Defaults to saving in place")
    args = parser.parse_args(argv)

    if args.manifest:
        if args.docx or args.image or args.anchor:
            raise SystemExit("--manifest cannot be combined with --docx/--image/--anchor")
        return run_manifest(args)
    return run_single(args)


if __name__ == "__main__":
    sys.exit(main())
