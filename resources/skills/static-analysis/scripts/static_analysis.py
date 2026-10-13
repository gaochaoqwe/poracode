#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
static-analysis 技能确定性辅助脚本（AgentEmployee 内置 Python 运行时，自带 openpyxl）。

子命令：
  parse-excel <xlsx> [--out issues.json]
      解析科代静态分析工具输出的 Excel：优先 GJB_8114 相关 sheet，
      按列映射（编号A/缺陷名称D/问题文件K/问题行数L/问题描述M）逐行提取问题，
      并从缺陷名称中提取 GJB 规则编号。输出 issues 数组（不读进对话，只取结果 JSON）。

  match --issues issues.json --rules-dir <dir> [--selected r1,r2...] [--out problems.json]
      对每条问题做规则匹配（先用缺陷名内嵌 GJB_R_... 规则号，否则按 keywords 命中
      gjb/cwe 规则库）、基础过滤（排除"建议"）、生成四段式问题单 + 测评单位意见，
      输出 problems（对齐 shared/staticAnalysisReport.ts 契约；AI 字段留空待 agent 判定）。

  emit-report --task-id <id> --mode <mode> [--selected r1,r2...]
      [--problems problems.json] [--files files.json] [--out report.json]
      组装 ingest 兼容的 report.json（schemaVersion/taskId/mode/files/problems/selectedRules）。
      --problems 省略则 problems=[]；--files 为 [{file_name,file_path}] 时脚本按路径读取
      source_code 与函数名片段。

  emit-progress --task-id <id> --mode <mode> --status running --progress 60
      [--total-files n --processed-files n --total-problems n --total-rules n
       --accurate-rules n --uncertain-rules n --current-rule-code r
       --current-analyzing-issue s --tool-calls n] [--out progress.json]
      写 progress.json（对齐契约）。

  rules --rules-dir <dir> [--type gjb|cwe] [--selected r1,r2...] [--out rules.json]
      导出规则（供纯AI模式 agent 查看选中的规则库条目）。
"""

import argparse
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path

# ---------- 科代 Excel 列映射（对齐原项目 filter/config/config.py） ----------
REQUIRED_COLUMNS = {
    "编号": 0,
    "缺陷名称": 3,
    "问题文件": 10,
    "问题行数": 11,
    "问题描述": 12,
}
TARGET_SHEET_PATTERN = "GJB_8114"
SKIP_SHEETS = ["项目详情", "你的问题汇总"]
EXCLUDE_KEYWORDS = ["建议"]
INCLUDE_KEYWORDS = [
    "不规范", "不清晰", "不符合", "不充分", "不一致", "混乱", "不对", "缺失", "缺少",
    "错误", "没有", "不全", "同名", "存在未初始化变量", "存在死循环", "存在空语句",
    "数组边界没有界定", "存在数组越界", "重定义", "不匹配", "无法", "未", "没", "数组越界",
]
GJB_CODE_RE = re.compile(r"GJB_R_(\d{2})_(\d{2})_(\d{2})")

RULES_DIR_CANDIDATES = [
    # 由技能注入时的 dir 相对推导：<skill_dir>/../../static-analysis/rules
    lambda skill_dir: (Path(skill_dir).resolve().parent.parent / "static-analysis" / "rules"),
    # 仓库内：<repo>/resources/static-analysis/rules
    lambda _: Path(__file__).resolve().parents[3] / "static-analysis" / "rules",
    # 打包后：process.resourcesPath/static-analysis/rules（由 agent 显式 --rules-dir 传入优先）
]


def find_rules_dir(explicit: str | None, skill_dir: str | None) -> Path | None:
    if explicit and Path(explicit).exists():
        return Path(explicit)
    cands = []
    if skill_dir:
        cands.append(RULES_DIR_CANDIDATES[0](skill_dir))
    cands.append(RULES_DIR_CANDIDATES[1](None))
    for c in cands:
        if c.exists():
            return c
    return None


def load_rules(rules_dir: Path) -> dict[str, list[dict]]:
    out = {"gjb": [], "cwe": [], "fpga": []}
    for typ, fname in (("gjb", "gjb8114.json"), ("cwe", "cwe.json")):
        p = rules_dir / fname
        if not p.exists():
            continue
        data = json.loads(p.read_text(encoding="utf-8-sig"))
        items = data if isinstance(data, list) else list(data.values())
        out[typ] = items
    return out


# ---------- 科代 Excel 解析 ----------

def _pick_process_sheets(sheetnames):
    targets = [s for s in sheetnames if TARGET_SHEET_PATTERN in s]
    if targets:
        return targets
    return [s for s in sheetnames if s not in SKIP_SHEETS]


def parse_excel(xlsx: Path):
    try:
        import openpyxl
    except Exception as e:  # pragma: no cover
        return {"success": False, "error": f"openpyxl 不可用: {e}"}
    wb = openpyxl.load_workbook(str(xlsx), read_only=True, data_only=True)
    try:
        sheets = _pick_process_sheets(wb.sheetnames)
        if not sheets:
            return {"success": False, "error": f"没有可处理的工作表。全部: {', '.join(wb.sheetnames)}"}
        issues = []
        for sn in sheets:
            ws = wb[sn]
            if ws.max_row < 2 or ws.max_column < 13:
                continue
            for row in ws.iter_rows(min_row=2, values_only=True):
                if len(row) <= 12:
                    continue
                defect = str(row[3]) if row[3] is not None else ""
                if not defect.strip():
                    continue
                issue_id = str(row[0]) if row[0] is not None else ""
                issue_file = str(row[10]) if row[10] is not None else ""
                issue_line = str(row[11]) if row[11] is not None else ""
                desc = str(row[12]) if row[12] is not None else ""
                desc = desc.replace("在(", "在").replace(")文件", "文件")
                desc = desc.replace("第(", "第").replace(")行", "行")
                desc = desc.replace("数字[", "数字").replace("件[", "件").replace("]件", "件")
                m = GJB_CODE_RE.search(defect)
                issues.append({
                    "issue_id": issue_id,
                    "issue_name": defect.strip(),
                    "gjb_rule_code": f"GJB_R_{m.group(1)}_{m.group(2)}_{m.group(3)}" if m else "",
                    "issue_file": issue_file.strip(),
                    "issue_line": issue_line.strip(),
                    "issue_description": desc.strip(),
                    "sheet_name": sn,
                    "row_index": row[0].row if hasattr(row[0], "row") else None,
                })
        return {"success": True, "total": len(issues), "sheets": sheets, "issues": issues}
    finally:
        wb.close()


# ---------- 规则匹配 + 四段式 ----------

def _convert_to_requirement(name: str) -> str:
    pairs = [
        ("不规范", "规范"), ("不清晰", "清晰"), ("不符合", "符合"), ("不充分", "充分"),
        ("不一致", "一致"), ("混乱", "清晰"), ("不对", "正确"), ("缺失", "完整"),
        ("缺少", "需要补充"), ("错误", "正确"), ("没有", "需要补充"), ("不全", "需要进行详细补充"),
        ("同名", "不能同名，具有各自独立的名称"),
        ("存在未初始化变量", "需要对变量进行初始化"),
        ("存在死循环", "不存在死循环"),
        ("存在空语句", "不存在空语句"),
        ("数组边界没有界定", "需要对数组边界进行界定"),
        ("存在数组越界", "不存在数组越界"),
        ("重定义", "不能重定义，需要使用不同的名称"),
        ("不匹配", "匹配"), ("无法", "能够"), ("未", "已"), ("没", "有"),
        ("数组越界", "数组不越界"),
    ]
    for old, new in pairs:
        if old in name:
            return re.sub(old, new, name)
    return name


def _impact_analysis(name: str) -> str:
    if "越界" in name:
        return "可能导致程序崩溃或产生不可预期的行为，严重影响软件稳定性"
    if "内存" in name and "泄漏" in name:
        return "可能导致系统内存不足，长时间运行后影响系统性能"
    if "死循环" in name:
        return "可能导致程序无响应，严重影响软件可用性"
    if "未初始化" in name:
        return "可能导致程序行为不确定，影响软件可靠性"
    if "编码规范" in name or "命名" in name:
        return "影响代码可读性和可维护性，不利于后续开发和维护工作"
    if "注释" in name:
        return "影响代码可理解性，不利于代码维护和团队协作"
    if "格式" in name:
        return "影响代码整体质量，不利于代码审查和维护"
    return f"{name},在某些特定场景下容易引发代码逻辑混乱，进而引发软件故障。"


def _four_part(name, source, desc, requirement) -> str:
    impact = _impact_analysis(name)
    text = (
        f"需求描述:{requirement}。\n"
        f"问题描述:{name}。\n"
        f"问题定位: {source},具体内容如下：\"{desc}\",如上述内容所示,{name}。\n"
        f"影响分析:{impact}"
    )
    return text.replace("。。", "。")


def _classify_source(name, source):
    if ".c" in source or ".h" in source or ".cpp" in source or "源代码" in source:
        typ, opinion = "程序问题", f"请开发人员针对{name}问题对代码进行修改，保证代码实现正确。"
    elif "文档" in source:
        typ, opinion = "软件需求问题", "请开发人员对文档问题进行修改,确保文档描述的正确且测试人员能够理解测试需求。"
    elif "任务书" in source:
        typ, opinion = "用户需求问题", "请开发人员对文档问题进行修改,确保文档描述的正确且测试人员能够理解测试需求。"
    elif ("规格说明" in source or "设计说明" in source or "影响域分析报告" in source or "接口" in source):
        typ, opinion = "软件需求问题", "请开发人员对文档问题进行修改,确保文档描述的正确且测试人员能够理解测试需求。"
    else:
        typ, opinion = "程序问题", f"请开发人员针对{name}问题对代码进行修改，保证代码实现正确。"
    return typ, opinion


def _match_rules(issue, rules: dict[str, list[dict]], selected: set[str]):
    """返回 (rule_type, rule_code, accuracy, hit_keyword)。"""
    name = issue["issue_name"] or ""
    desc = issue["issue_description"] or ""
    blob = f"{name}\n{desc}"

    # 1) 缺陷名内嵌 GJB_R_... 规则号（科代原始格式）——转成规则库 key 未命中时按关键词找
    embedded = issue.get("gjb_rule_code", "")
    if embedded:
        # 本项目规则库 rule_code 形如 R-1-1-1，科代格式为 GJB_R_01_01_01
        for typ in ("gjb", "cwe", "fpga"):
            for rule in rules[typ]:
                if rule.get("rule_code") and embedded.replace("GJB_R_", "R-").replace("_", "-") == rule["rule_code"]:
                    if not selected or rule["rule_code"] in selected:
                        return typ, rule["rule_code"], rule.get("accuracy_type", "uncertain"), embedded
    # 2) 关键词命中
    for typ in ("gjb", "cwe", "fpga"):
        for rule in rules[typ]:
            if selected and rule.get("rule_code") not in selected:
                continue
            kws = rule.get("keywords") or []
            for kw in kws:
                k = str(kw.get("keyword", ""))
                if k and k in blob:
                    return typ, rule.get("rule_code", ""), rule.get("accuracy_type", "uncertain"), k
    return None, "", "uncertain", ""


def _build_problem(issue, matched, level) -> dict:
    name = issue["issue_name"] or ""
    file = issue.get("issue_file") or ""
    line = issue.get("issue_line") or ""
    source = f"{file}文件{line}行" if file else "源代码"
    if "、" in source:
        count = source.count("、") + 1
    else:
        count = 1
    typ, opinion = _classify_source(name, source)
    requirement = _convert_to_requirement(name)
    four = _four_part(name, source, issue.get("issue_description") or "", requirement)
    # 拆分四段（供渲染端 description_requirement 等字段）
    req_part = issue.get("issue_description") or ""
    return {
        "problemId": f"SA-{issue.get('row_index') or 'row'}-{issue.get('issue_id') or name[:8]}",
        "issue_name": name,
        "issue_file": file,
        "issue_line": int(issue["issue_line"]) if (issue.get("issue_line") or "").isdigit() else None,
        "issue_location": source,
        "issue_description": issue.get("issue_description") or "",
        "issue_level": level,
        "rule_code": matched[1],
        "rule_type": matched[0],
        "rule_accuracy_type": matched[2],
        "ai_analyzed": False,
        "ai_is_real_issue": None,
        "ai_confidence": None,
        "ai_analysis_result": "",
        "ai_reasoning": "",
        "ai_recommendations": "",
        "problem_status": 0,
        "is_filtered": False,
        "filter_reason": "",
        "file_name": file,
        "violation_rules": matched[1],
        "description_requirement": requirement,
        "description_issue": name,
        "description_location": source,
        "description_impact": _impact_analysis(name),
        "problem_sheet": {
            "issue_name": name,
            "four_part_description": four,
            "test_unit_opinion": opinion,
        },
        "function_name": "",
        "source_code": "",
        "sheet_name": issue.get("sheet_name", ""),
        "row_index": issue.get("row_index"),
    }


def match_issues(issues, rules, selected):
    problems = []
    filtered = []
    for issue in issues:
        name = issue.get("issue_name") or ""
        if not name.strip():
            continue
        if any(k in (issue.get("issue_description") or "") for k in EXCLUDE_KEYWORDS):
            issue = dict(issue); issue["_filtered"] = "含排除关键词"
        matched = _match_rules(issue, rules, selected)
        if any(k in name or k in (issue.get("issue_description") or "") for k in EXCLUDE_KEYWORDS):
            p = _build_problem(issue, ("", "", "uncertain", ""), "一般问题")
            p["is_filtered"] = True
            p["filter_reason"] = "含排除关键词(建议)"
            p["ai_analyzed"] = True
            p["ai_is_real_issue"] = False
            filtered.append(p)
            continue
        level = "重要问题" if any(x in name for x in ("越界", "章节缺失", "失效")) else "一般问题"
        problems.append(_build_problem(issue, matched, level))
    return problems, filtered


# ---------- 报告组装 ----------

def _read_code_snippets(files):
    out = []
    for f in files or []:
        path = f.get("file_path") or f.get("path") or ""
        name = f.get("file_name") or (path.split("/")[-1].split("\\")[-1] if path else "")
        if not path:
            continue
        snippet = ""
        func = ""
        try:
            p = Path(path)
            if p.exists() and p.is_file():
                snippet = p.read_text(encoding="utf-8", errors="replace")[:8000]
                # 尝试提取第一个函数名
                m = re.search(r"(?:^|\n)\s*(?:static\s+)?(?:[\w:<>*&]+\s+)+([A-Za-z_]\w*)\s*\(", snippet)
                if m:
                    func = m.group(1)
        except Exception:
            snippet = ""
        out.append({"file_name": name, "file_path": path, "source_code": snippet})
        if func:
            out[-1]["function_name"] = func
    return out


def main() -> int:
    ap = argparse.ArgumentParser(prog="static_analysis.py")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("parse-excel")
    p.add_argument("xlsx")
    p.add_argument("--out")

    m = sub.add_parser("match")
    m.add_argument("--issues", required=True)
    m.add_argument("--rules-dir")
    m.add_argument("--skill-dir")
    m.add_argument("--selected", default="")
    m.add_argument("--out")

    e = sub.add_parser("emit-report")
    e.add_argument("--task-id", required=True)
    e.add_argument("--mode", required=True)
    e.add_argument("--selected", default="")
    e.add_argument("--problems")
    e.add_argument("--files")
    e.add_argument("--out", required=True)

    pr = sub.add_parser("emit-progress")
    pr.add_argument("--task-id", required=True)
    pr.add_argument("--mode", required=True)
    pr.add_argument("--status", default="running")
    pr.add_argument("--progress", type=int, default=0)
    pr.add_argument("--total-files", type=int, default=0)
    pr.add_argument("--processed-files", type=int, default=0)
    pr.add_argument("--total-problems", type=int, default=0)
    pr.add_argument("--total-rules", type=int, default=0)
    pr.add_argument("--accurate-rules", type=int, default=0)
    pr.add_argument("--uncertain-rules", type=int, default=0)
    pr.add_argument("--current-rule-code", default="")
    pr.add_argument("--current-analyzing-issue", default="")
    pr.add_argument("--tool-calls", type=int, default=0)
    pr.add_argument("--out", required=True)

    r = sub.add_parser("rules")
    r.add_argument("--rules-dir")
    r.add_argument("--skill-dir")
    r.add_argument("--type", default="")
    r.add_argument("--selected", default="")
    r.add_argument("--out")

    args = ap.parse_args()

    if args.cmd == "parse-excel":
        res = parse_excel(Path(args.xlsx))
        if args.out:
            Path(args.out).write_text(json.dumps(res, ensure_ascii=False, indent=2), encoding="utf-8")
        else:
            print(json.dumps(res, ensure_ascii=False))
        return 0 if res.get("success") else 1

    if args.cmd in ("match", "rules"):
        rd = find_rules_dir(args.rules_dir, getattr(args, "skill_dir", None))
        if not rd:
            print(json.dumps({"success": False, "error": "未找到规则库目录"}, ensure_ascii=False))
            return 1
        rules = load_rules(rd)
        selected = {s for s in (args.selected or "").split(",") if s} if getattr(args, "selected", "") else set()

    if args.cmd == "rules":
        wanted = args.type.lower()
        types = ["gjb", "cwe", "fpga"] if not wanted else [wanted]
        out = []
        for t in types:
            for rule in rules[t]:
                if selected and rule.get("rule_code") not in selected:
                    continue
                out.append({**rule, "rule_type": t})
        payload = {"success": True, "rules": out, "total": len(out)}
        if args.out:
            Path(args.out).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        else:
            print(json.dumps(payload, ensure_ascii=False))
        return 0

    if args.cmd == "match":
        issues = json.loads(Path(args.issues).read_text(encoding="utf-8-sig"))
        if isinstance(issues, dict):
            issues = issues.get("issues", [])
        problems, filtered = match_issues(issues, rules, selected)
        # 被基础过滤丢弃的问题也进 problems，标记 is_filtered，供渲染端显示过滤状态
        problems = problems + filtered
        payload = {"success": True, "total": len(problems), "problems": problems}
        if args.out:
            Path(args.out).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        else:
            print(json.dumps(payload, ensure_ascii=False))
        return 0

    if args.cmd == "emit-report":
        problems = []
        if args.problems and Path(args.problems).exists():
            data = json.loads(Path(args.problems).read_text(encoding="utf-8-sig"))
            problems = data.get("problems", data) if isinstance(data, dict) else data
        files = []
        if args.files and Path(args.files).exists():
            data = json.loads(Path(args.files).read_text(encoding="utf-8-sig"))
            files = _read_code_snippets(data.get("files", data) if isinstance(data, dict) else data)
        selected = [s for s in (args.selected or "").split(",") if s]
        report = {
            "schemaVersion": 1,
            "taskId": args.task_id,
            "mode": args.mode,
            "fileName": f"{args.task_id}-{args.mode}-report.json",
            "files": files,
            "problems": problems,
            "selectedRules": selected,
        }
        Path(args.out).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return 0

    if args.cmd == "emit-progress":
        prog = {
            "taskId": args.task_id,
            "mode": args.mode,
            "status": args.status,
            "progress": args.progress,
            "totalFiles": args.total_files,
            "processedFiles": args.processed_files,
            "totalProblems": args.total_problems,
            "totalRules": args.total_rules,
            "accurateRules": args.accurate_rules,
            "uncertainRules": args.uncertain_rules,
            "totalUnits": 0,
            "processedUnits": 0,
            "currentRuleCode": args.current_rule_code or None,
            "currentAnalyzingIssue": args.current_analyzing_issue or None,
            "toolCallsCount": args.tool_calls,
            "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
        Path(args.out).write_text(json.dumps(prog, ensure_ascii=False, indent=2), encoding="utf-8")
        return 0

    ap.print_help()
    return 1


if __name__ == "__main__":
    sys.exit(main())
