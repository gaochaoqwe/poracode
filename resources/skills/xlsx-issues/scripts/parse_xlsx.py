#!/usr/bin/env python3
"""Read-only inspector for .xlsx workbooks, using the bundled openpyxl.

The agent cannot unzip OOXML itself, so it reads a workbook through this
script: first `--list` to see every sheet's shape and header, then `--sheet`
to dump the issue sheet's rows. Only `.xlsx` is supported (legacy `.xls` is a
different binary format).

Usage:
  parse_xlsx.py --list <file.xlsx> [--preview-rows N] [--out out.json]
  parse_xlsx.py --sheet <name|index> <file.xlsx> [--out out.json]
"""

import argparse
import json
import sys

try:
    import openpyxl
except ImportError:  # pragma: no cover - depends on the interpreter used
    sys.exit("openpyxl is not available in this interpreter; use the bundled Python.")

# Emit UTF-8 regardless of the Windows console codepage, so the agent reads
# valid JSON with intact non-ASCII text.
for stream in (sys.stdout, sys.stderr):
    try:
        stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):  # pragma: no cover - non-conforming stream
        pass


def to_json_value(value):
    """JSON-safe scalar: openpyxl may return dates, decimals, etc."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def resolve_sheet(workbook, ref):
    if ref in workbook.sheetnames:
        return workbook[ref]
    try:
        return workbook.worksheets[int(ref)]
    except (ValueError, IndexError):
        sys.exit("No sheet named or indexed '%s'. Sheets: %s" % (ref, ", ".join(workbook.sheetnames)))


def sheet_info(worksheet, preview_rows):
    header = []
    preview = []
    for index, row in enumerate(worksheet.iter_rows(values_only=True)):
        values = [to_json_value(cell) for cell in row]
        if index == 0:
            header = values
        if index < preview_rows:
            preview.append(values)
        else:
            break
    return {
        "name": worksheet.title,
        "maxRow": worksheet.max_row,
        "maxColumn": worksheet.max_column,
        "header": header,
        "preview": preview,
    }


def main():
    parser = argparse.ArgumentParser(description="Inspect an .xlsx workbook (read-only).")
    parser.add_argument("file", help="path to the .xlsx workbook")
    parser.add_argument("--list", action="store_true", help="list all sheets with dimensions and a preview")
    parser.add_argument("--sheet", help="dump one sheet by name or 0-based index")
    parser.add_argument("--preview-rows", type=int, default=5, help="rows to preview per sheet in --list (default 5)")
    parser.add_argument("--out", help="write JSON to this file instead of stdout")
    args = parser.parse_args()

    if not args.file.lower().endswith(".xlsx"):
        sys.exit("Not an .xlsx file (legacy .xls is not supported): %s" % args.file)
    if not args.list and args.sheet is None:
        parser.error("choose --list or --sheet")

    workbook = openpyxl.load_workbook(args.file, read_only=True, data_only=True)
    try:
        if args.list:
            payload = {
                "file": args.file,
                "sheets": [sheet_info(ws, args.preview_rows) for ws in workbook.worksheets],
            }
        else:
            worksheet = resolve_sheet(workbook, args.sheet)
            rows = [
                [to_json_value(cell) for cell in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            payload = {"file": args.file, "sheet": worksheet.title, "rows": rows}
    finally:
        workbook.close()

    text = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as handle:
            handle.write(text)
        print("wrote %s" % args.out)
    else:
        print(text)


if __name__ == "__main__":
    main()
