#!/usr/bin/env python3
"""Write per-row fix status back into a COPY of an .xlsx issue workbook.

The original workbook is never modified: this loads it, appends (or reuses)
status/note columns, writes one entry per issue row, and saves to `--out`.
Refuses when `--out` equals `--file`.

Usage:
  write_status.py --file src.xlsx --out dst.xlsx --sheet <name|index> \
      --statuses statuses.json [--status-header 修复状态] [--note-header 修复说明] [--merge]

statuses.json:
  { "rows": [ { "row": 5, "status": "已修复", "note": "src/a.ts:login —— 去掉了未用变量" } ] }
  `row` is 1-based and matches the spreadsheet row (header row is 1).

--merge: when `--out` already exists, load it as the base and update only the
given rows (accumulate progress across batches) instead of rebuilding from the
source workbook. Without --merge, or when `--out` is absent, the source is the
base. The source workbook is never written.
"""

import argparse
import json
import os
import sys

try:
    import openpyxl
except ImportError:  # pragma: no cover - depends on the interpreter used
    sys.exit("openpyxl is not available in this interpreter; use the bundled Python.")

# Emit UTF-8 regardless of the Windows console codepage.
for stream in (sys.stdout, sys.stderr):
    try:
        stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):  # pragma: no cover - non-conforming stream
        pass


def resolve_sheet(workbook, ref):
    if ref in workbook.sheetnames:
        return workbook[ref]
    try:
        return workbook.worksheets[int(ref)]
    except (ValueError, IndexError):
        sys.exit("No sheet named or indexed '%s'. Sheets: %s" % (ref, ", ".join(workbook.sheetnames)))


def find_header_column(worksheet, header_text):
    """Return the 1-based column whose row-1 header equals `header_text`, else None."""
    target = header_text.strip()
    for column in range(1, worksheet.max_column + 1):
        value = worksheet.cell(row=1, column=column).value
        if value is not None and str(value).strip() == target:
            return column
    return None


def ensure_column(worksheet, header_text, after_column):
    column = find_header_column(worksheet, header_text)
    if column is not None:
        return column
    column = max(worksheet.max_column, after_column) + 1
    worksheet.cell(row=1, column=column, value=header_text)
    return column


def main():
    parser = argparse.ArgumentParser(description="Write per-row status into a copy of an .xlsx workbook.")
    parser.add_argument("--file", required=True, help="source .xlsx workbook (read-only)")
    parser.add_argument("--out", required=True, help="destination .xlsx workbook (must differ from --file)")
    parser.add_argument("--sheet", required=True, help="sheet name or 0-based index holding the issue rows")
    parser.add_argument("--statuses", required=True, help="JSON file with a top-level 'rows' array")
    parser.add_argument("--status-header", default="修复状态", help="status column header (created if absent)")
    parser.add_argument("--note-header", default="修复说明", help="note column header (created if absent)")
    parser.add_argument("--merge", action="store_true",
                        help="update the existing --out workbook in place (accumulate across batches)")
    args = parser.parse_args()

    if os.path.abspath(args.file) == os.path.abspath(args.out):
        sys.exit("refusing to overwrite the source workbook; choose a different --out")

    with open(args.statuses, encoding="utf-8") as handle:
        data = json.load(handle)
    entries = data.get("rows", [])
    if not isinstance(entries, list):
        sys.exit("'rows' must be a JSON array")

    # With --merge, keep the copy already under --out (it carries earlier
    # batches' statuses); otherwise rebuild from the untouched source.
    base = args.out if (args.merge and os.path.exists(args.out)) else args.file
    workbook = openpyxl.load_workbook(base)
    try:
        worksheet = resolve_sheet(workbook, args.sheet)
        status_column = ensure_column(worksheet, args.status_header, 0)
        note_column = ensure_column(worksheet, args.note_header, status_column)

        for entry in entries:
            row = entry.get("row")
            if not isinstance(row, int) or row < 1:
                sys.exit("each status entry needs a 1-based integer 'row': %r" % (entry,))
            worksheet.cell(row=row, column=status_column, value=entry.get("status"))
            if entry.get("note") is not None:
                worksheet.cell(row=row, column=note_column, value=entry["note"])
    finally:
        pass

    workbook.save(args.out)
    workbook.close()
    print("wrote %s (%d rows)" % (args.out, len(entries)))


if __name__ == "__main__":
    main()
