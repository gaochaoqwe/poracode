# -*- coding: utf-8 -*-
"""将正向追踪表模板保留的尾部两行提示文本改写为真实追踪数据（只改内容，不增删行）。"""
import json
from pathlib import Path

from docx import Document

DOC = Path(r"E:\code\AgentEmployee-master\docs\雷达仿真系统软件需求规格说明-GJB438C-2021.docx")
ROWS_FILE = Path(r"E:\code\AgentEmployee-master\docs\研制任务书\srs\forward_rows.json")
LOC0 = ("《XXX课题可行性研究报告》中的指标", "《XXX课题可行性研究报告》中的指标", "软件需求", "软件需求")
LOC1 = ("指标内容", "章节号", "需求名称", "章节号")
KEYS = ["source_indicator_content", "source_chapter_number", "requirement_name", "requirement_chapter_number"]


def row_texts(row):
    return tuple(c.text.strip() for c in row.cells)


def find_table(doc):
    hits = [t for t in doc.tables
            if len(t.rows) >= 2 and row_texts(t.rows[0]) == LOC0 and row_texts(t.rows[1]) == LOC1]
    assert len(hits) == 1, "locate forward traceability table: %d" % len(hits)
    return hits[0]


def set_cell(cell, text):
    first = cell.paragraphs[0]
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)
    for run in list(first.runs):
        run.text = ""
    if first.runs:
        first.runs[0].text = text
    else:
        first.add_run(text)


def main():
    rows = json.loads(ROWS_FILE.read_text(encoding="utf-8"))
    doc = Document(str(DOC))
    table = find_table(doc)
    for row, data in zip(table.rows[-2:], rows[-2:]):
        for cell, key in zip(row.cells, KEYS):
            set_cell(cell, data[key])
    doc.save(str(DOC))
    print("tail rows filled:", [r["source_chapter_number"] for r in rows[-2:]])
    print("table rows now:", len(table.rows))


if __name__ == "__main__":
    main()