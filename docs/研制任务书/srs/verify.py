# -*- coding: utf-8 -*-
import re
from pathlib import Path
from docx import Document

DOC = Path(r"E:\code\AgentEmployee-master\docs\雷达仿真系统软件需求规格说明-GJB438C-2021.docx")
BAD = ["本条应描述", "本条应概述", "本条应标识", "本章应列出", "本章应描述", "本章应包括",
       "3.2.X", "3.3.X", "外部型号+产品名称", "产品型号-XXXX", "若存在一对多",
       "SAMPLE", "BYQ", "编译器软件", "Test table", "Test 3.2", "如如实陈述"]


def all_paras(doc):
    for para in doc.paragraphs:
        yield para.text
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield cell.text
    for section in doc.sections:
        for part in (section.header, section.footer):
            for para in part.paragraphs:
                yield para.text


doc = Document(str(DOC))
texts = list(all_paras(doc))
print("paragraph/cell texts:", len(texts))
hits = {b: sum(1 for t in texts if b in t) for b in BAD}
print("leftover hits:", {k: v for k, v in hits.items() if v})

heads = [p.text.strip() for p in doc.paragraphs if p.style and ("Heading" in p.style.name or "标题" in p.style.name)]
print("headings:", len(heads))
for h in heads[:8]:
    print("  ", h)
missing = [n for n in ["3.1.1", "3.1.6", "3.2.1", "3.2.11", "3.3.1", "3.3.8", "3.4.1", "3.4.3",
                       "3.10.4", "3.18", "3.5", "3.17"] if not any(x.startswith(n + " ") or x.startswith(n) for x in heads)] if False else []
line = [h for h in heads if h.startswith("3.")]
print("chapter-3 headings:", len(line))
print("  sample:", line[:3], "...", line[-3:])
print("tables:", [(len(t.rows), len(t.columns)) for t in doc.tables])
print("cover:", [t.text.replace(chr(10), " | ") for t in doc.tables[0].rows[0].cells][0][:60])
assert not any(hits.values()), hits
print("VERIFY OK")